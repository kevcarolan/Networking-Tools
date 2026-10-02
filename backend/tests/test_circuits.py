import io

import pytest
from openpyxl import Workbook

from app.tools.circuits.importer import (CircuitImportError, match_switches, normalize_mac,
                                         normalize_port, parse_workbook, switch_matches)

BANDS = ["SERVICE", "", "DEVICE INFORMATION", "", "", "", "", "LOCATION", "", "IP DETAILS", "",
         "", "NETWORK PORT", "", "", "", "DRAWING DETAILS", ""]
HEADER = ["IP Partition", "Service", "ENGineering prefix", "Device Name", "Device Type",
          "Manufacturer", "Hardware Mac Address", "Floor Level", "Room / Area", "IP Address",
          "Subnet Mask", "Default Gateway", "VLAN", "TYPE", "PORT", "DEVICE", "Drawing Number",
          "Device Drawing Reference"]
ROWS = [
    ["TSP", "ACS", None, "A-DUB01-GH-100-3", "eDCM 400", "CEM", "00304604c02f", "Ground Floor",
     "SEC01", "10.3.34.62", "255.255.254.0", "10.3.34.1", 34, "STATIC", "Gi2/0/13", "GH-AS01",
     "DUB01-TSP-GH-00-DR-SEC-68110", "Intelligent Door Controller"],
    ["TSP", "ACS", None, "A-DUB01-GH-100-2", "eDCM 400", "CEM", "00304604c05e", "Ground Floor",
     "SEC01", "10.3.34.61", "255.255.254.0", "10.3.34.1", 34, "STATIC", "Gi2/0/12", "GH-AS01",
     "DUB01-TSP-GH-00-DR-SEC-68110", "Intelligent Door Controller"],
    ["SE BMS", "BMS", "DUB01_DC1_00_N_COR11_LL_B_BMS_CP03", "LL_B_BMS_CP03", None, None,
     "0050062B2269", None, None, "10.3.66.16", "255.255.255.0", "10.3.66.1", 66, "STATIC",
     "Gi1/0/3", "GH-AS01", None, None],
    ["TSP", "Intercom", None, "IC-DUB01-SW-009", "2N IP Force 1button+Hd", "Axis 2N",
     "7C1EB3066248", "N/A", "Vehicle Gate", "10.3.36.34", "255.255.255.0", "10.3.36.1", 36,
     "STATIC", "Gi2/0/8", "GH-AS02", "DUB01-TSP-EP-XX-DR-SEC-68111", "IC-DUB01-SW-009"],
    ["TSP", "Intercom", None, "IC-DUB01-GH-0004", "2N IP Force 1button+Hd", "Axis 2N",
     "7C1EB3066292", "Ground Floor", "Turnstile", "10.3.36.39", "255.255.255.0", "10.3.36.1",
     36, "STATIC", "Gi2/0/15", "GH-AS01", None, "IC-DUB01-GH-0004"],
]


def make_xlsx(rows=ROWS, header=HEADER, bands=BANDS, extra_sheet_first=False) -> bytes:
    wb = Workbook()
    ws = wb.active
    if extra_sheet_first:
        ws.title = "Notes"
        ws.append(["This sheet has no header"])
        ws = wb.create_sheet("Circuits")
    else:
        ws.title = "Circuits"
    ws.append([None] * len(header))  # row 1: empty
    ws.append(bands)                 # row 2: band names
    ws.append(header)                # row 3: column names
    for r in rows:
        ws.append(r)
    ws.append([None] * len(header))  # trailing empty row is ignored
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- parsing and normalisation ----------------------------------------------

def test_parse_finds_header_under_bands():
    parsed = parse_workbook(make_xlsx(extra_sheet_first=True))
    assert parsed.sheet == "Circuits" and parsed.header_row == 3
    assert len(parsed.rows) == 5 and parsed.rows[0].row_no == 4
    r = parsed.rows[0]
    assert r.values["device_name"] == "A-DUB01-GH-100-3"
    assert r.values["vlan"] == "34" and r.values["switch"] == "GH-AS01"
    assert r.port_norm == "gi2/0/13" and r.switch_norm == "gh-as01"
    assert r.mac_norm == "00:30:46:04:c0:2f" and r.warnings == []
    assert parsed.rows[2].values["eng_prefix"] == "DUB01_DC1_00_N_COR11_LL_B_BMS_CP03"


def test_parse_column_order_does_not_matter():
    order = list(reversed(range(len(HEADER))))
    parsed = parse_workbook(make_xlsx(rows=[[r[i] for i in order] for r in ROWS],
                                      header=[HEADER[i] for i in order], bands=[]))
    assert parsed.rows[3].values["switch"] == "GH-AS02"
    assert parsed.rows[3].values["service"] == "Intercom"


def test_parse_warnings():
    bad = list(ROWS[0])
    bad[9], bad[10], bad[12], bad[6], bad[14] = "10.3.34.999", "255.0.255.0", 5000, "xyz", None
    parsed = parse_workbook(make_xlsx(rows=[bad]))
    w = " | ".join(parsed.rows[0].warnings)
    for text in ("no switch/port", "IP address", "subnet mask", "VLAN", "MAC"):
        assert text in w


def test_parse_rejects_bad_files():
    with pytest.raises(CircuitImportError):
        parse_workbook(b"not an excel file")
    with pytest.raises(CircuitImportError, match="DEVICE"):
        parse_workbook(make_xlsx(header=[h if h != "DEVICE" else "Switch name" for h in HEADER]))
    wb = Workbook()
    wb.active.append(["a", "b"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(CircuitImportError, match="header"):
        parse_workbook(buf.getvalue())


@pytest.mark.parametrize("raw, norm", [
    ("Gi2/0/13", "gi2/0/13"), ("GigabitEthernet2/0/13", "gi2/0/13"), ("gi 2/0/13", "gi2/0/13"),
    ("Te1/1/1", "te1/1/1"), ("TenGigabitEthernet1/1/1", "te1/1/1"), ("Fa0/1", "fa0/1"),
    ("Eth1/10", "eth1/10"), ("Ethernet1/10", "eth1/10"), ("port1.0.1", "port1.0.1"),
    ("Po10", "po10"), ("Port-channel10", "po10"),
])
def test_normalize_port(raw, norm):
    assert normalize_port(raw) == norm


def test_normalize_mac():
    assert normalize_mac("0050062B2269") == "00:50:06:2b:22:69"
    assert normalize_mac("00-50-06-2B-22-69") == normalize_mac("0050.062b.2269")
    assert normalize_mac("12345") == ""


def test_switch_matching():
    assert switch_matches("GH-AS01", "gh-as01")
    assert switch_matches("gh-as01.corp.local", "gh-as01")
    assert switch_matches("DUB01-GH-AS01", "gh-as01")
    assert not switch_matches("GH-AS011", "gh-as01")
    names = {1: "DUB01-GH-AS01", 2: "GH-AS02", 3: "DUB02-GH-AS02", 4: "core1"}
    m = match_switches(names, {"gh-as01", "gh-as02", "gh-as09"})
    assert m["gh-as01"] == ("matched", [1])
    assert m["gh-as02"] == ("matched", [2])  # exact name wins over the suffix match
    assert m["gh-as09"] == ("unmatched", [])
    assert match_switches({1: "A-SW1", 2: "B-SW1"}, {"sw1"})["sw1"][0] == "ambiguous"


# --- API ----------------------------------------------------------------------

def _upload(client, data, filename="circuits.xlsx"):
    return client.put("/api/circuits/imports/upload", params={"filename": filename},
                      content=data)


@pytest.fixture
def switch(admin, credential):
    r = admin.post("/api/devices", json={"name": "DUB01-GH-AS01", "address": "10.3.0.11",
                                         "platform": "cisco_ios", "site": "DUB01",
                                         "credential_id": credential["id"]})
    assert r.status_code == 201, r.text
    return r.json()


def test_upload_preview_commit_and_query(admin, switch):
    assert admin.get("/api/circuits").json() == {"import": None, "circuits": []}
    r = _upload(admin, make_xlsx())
    assert r.status_code == 201, r.text
    imp = r.json()
    s = imp["summary"]
    assert imp["status"] == "pending" and imp["row_count"] == 5
    assert s["header_row"] == 3 and s["unmatched_switches"] == ["gh-as02"]
    assert s["rows_unmatched"] == 1 and s["diff"]["added"] == 5
    assert admin.get("/api/circuits").json()["circuits"] == []  # nothing live until commit

    assert admin.post(f"/api/circuits/imports/{imp['id']}/commit").status_code == 200
    data = admin.get("/api/circuits").json()
    assert data["import"]["status"] == "active" and len(data["circuits"]) == 5
    first = data["circuits"][0]
    assert first["match"] == "matched" and first["switch_device_name"] == "DUB01-GH-AS01"

    fd = admin.get(f"/api/circuits/for-device/{switch['id']}").json()
    assert fd["count"] == 4
    assert [(g["service"], g["vlan"], len(g["circuits"])) for g in fd["groups"]] == [
        ("ACS", "34", 2), ("BMS", "66", 1), ("Intercom", "36", 1)]

    csv = admin.get("/api/circuits/export.csv").text
    assert csv.splitlines()[0].startswith("IP Partition,Service,Engineering prefix")
    assert "A-DUB01-GH-100-3" in csv and "DUB01-GH-AS01" in csv


def test_reimport_shows_changes_and_keeps_history(admin, switch):
    first = _upload(admin, make_xlsx()).json()
    admin.post(f"/api/circuits/imports/{first['id']}/commit")

    rows = [list(r) for r in ROWS]
    rows[0][14] = "Gi2/0/20"          # moved port
    del rows[1]                       # removed
    rows.append(["TSP", "CCTV", None, "CAM-01", "Camera", "Axis", "ACCC8E000001", "Ground Floor",
                 "Lobby", "10.3.40.10", "255.255.255.0", "10.3.40.1", 40, "STATIC", "Gi1/0/5",
                 "GH-AS01", None, None])  # added
    second = _upload(admin, make_xlsx(rows=rows)).json()
    d = second["summary"]["diff"]
    assert (d["added"], d["removed"], d["changed"]) == (1, 1, 1)
    assert d["changed_sample"][0]["fields"] == ["port"]
    assert d["changed_sample"][0]["before"]["port"] == "Gi2/0/13"
    assert second["summary"]["replaces_import"] == first["id"]

    admin.post(f"/api/circuits/imports/{second['id']}/commit")
    history = admin.get("/api/circuits/imports").json()
    assert [i["status"] for i in history] == ["active", "superseded"]
    assert len(admin.get("/api/circuits").json()["circuits"]) == 5


def test_discard_and_replace_pending(admin):
    a = _upload(admin, make_xlsx()).json()
    b = _upload(admin, make_xlsx()).json()  # a new upload replaces the pending preview
    assert admin.get(f"/api/circuits/imports/{a['id']}").status_code == 404
    assert admin.delete(f"/api/circuits/imports/{b['id']}").status_code == 204
    assert admin.get("/api/circuits/imports").json() == []
    c = _upload(admin, make_xlsx()).json()
    admin.post(f"/api/circuits/imports/{c['id']}/commit")
    assert admin.delete(f"/api/circuits/imports/{c['id']}").status_code == 409


def test_upload_errors(admin):
    assert _upload(admin, make_xlsx(), filename="list.csv").status_code == 422
    r = _upload(admin, b"garbage")
    assert r.status_code == 422 and "xlsx" in r.json()["detail"]


def test_viewer_can_read_but_not_import(client, app, admin, switch):
    imp = _upload(admin, make_xlsx()).json()
    admin.post(f"/api/circuits/imports/{imp['id']}/commit")
    from app.core.auth import VIEWER, User

    app.state.authenticator.authenticate = lambda u, p: User(username=u, role=VIEWER)
    client.post("/api/auth/logout")
    client.post("/api/auth/login", json={"username": "v", "password": "x"})
    assert len(client.get("/api/circuits").json()["circuits"]) == 5
    assert client.get(f"/api/circuits/for-device/{switch['id']}").status_code == 200
    assert _upload(client, make_xlsx()).status_code == 403
    assert client.post(f"/api/circuits/imports/{imp['id']}/commit").status_code == 403
